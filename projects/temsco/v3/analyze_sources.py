from pathlib import Path
import json, hashlib, unicodedata
import openpyxl
from docx import Document
from pypdf import PdfReader

root = Path(__file__).parent
out = root / 'analysis'
out.mkdir(exist_ok=True)
manifest = []
for i, path in enumerate(sorted((root / 'source-materials').iterdir()), 1):
    if not path.is_file() or path.name.startswith('.'):
        continue
    name = unicodedata.normalize('NFC', path.name)
    item = {'id': f'S{i:02}', 'name': name, 'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
    lines = [name]
    if path.suffix == '.xlsx':
        wb = openpyxl.load_workbook(path, data_only=False)
        values = openpyxl.load_workbook(path, data_only=True)
        sheets = []
        for ws in wb:
            cells = []
            lines.append(f'\n## {ws.title} ({ws.sheet_state}) {ws.max_row} rows x {ws.max_column} columns')
            for row in ws:
                parts = []
                for c in row:
                    if c.value is None: continue
                    v = values[ws.title][c.coordinate].value
                    cells.append({'cell':c.coordinate,'value':v,'formula':c.value if c.data_type=='f' else None, 'format':c.number_format})
                    parts.append(f'{c.coordinate}: {v}' + (f' [={c.value[1:]}]' if c.data_type=='f' else ''))
                if parts: lines.append(' | '.join(parts))
            sheets.append({'name':ws.title,'state':ws.sheet_state,'cells':cells})
        (out / f'{item["id"]}.json').write_text(json.dumps(sheets, ensure_ascii=False, indent=2, default=str))
        item['sheets'] = [s['name'] for s in sheets]
    elif path.suffix == '.docx':
        doc = Document(path)
        for p in doc.paragraphs:
            if p.text: lines.append(p.text)
        for j, table in enumerate(doc.tables):
            lines.append(f'TABLE {j+1}')
            lines.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
    elif path.suffix == '.pdf':
        doc = PdfReader(path).pages
        item['pages'] = len(doc)
        for j, page in enumerate(doc):
            lines.append(f'\n## PAGE {j+1}\n'+page.extract_text())
        item['page_text_lengths'] = [len(p.extract_text()) for p in doc]
    elif path.suffix == '.txt':
        lines.append(path.read_text())
    (out / f'{item["id"]}.txt').write_text('\n'.join(lines))
    manifest.append(item)
(out / 'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps(manifest,ensure_ascii=False,indent=2))
